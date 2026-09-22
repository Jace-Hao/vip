# -*- coding: utf-8 -*-
"""
洗衣管家会员导出 · Excel 生成（make_xlsx.py）v1.1
由导出脚本生成的 JSON 数据文件生成排版好的 .xlsx：
  1) 会员导出_xxx.xlsx       常规表（9 列）
  2) 会员全量数据_xxx.xlsx   全量字段表
  3) 会员详情_xxx.xlsx       详情总表 + 订单明细 + 卡券明细（深度模式时）
依赖：openpyxl（本机已安装）
"""
import json
import os
import sys


def style_header(ws, ncols):
    from openpyxl.styles import Font, Alignment, PatternFill
    bold = Font(bold=True, color='FFFFFF')
    fill = PatternFill('solid', fgColor='4472C4')
    center = Alignment(horizontal='center', vertical='center')
    for c in ws[1]:
        c.font = bold
        c.fill = fill
        c.alignment = center
    ws.row_dimensions[1].height = 20


def col_letter(n):
    from openpyxl.utils import get_column_letter
    return get_column_letter(n)


def build_summary_xlsx(data, out_path):
    from openpyxl import Workbook
    headers = data.get('headers') or []
    rows = data.get('rows') or []
    wb = Workbook()
    ws = wb.active
    ws.title = '会员数据'
    ws.append(headers)
    style_header(ws, len(headers))
    for r in rows:
        ws.append(r)
    for col in (2, 5):
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            for c in row:
                if c.value is not None:
                    c.value = str(c.value)
                c.number_format = '@'
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        for c in row:
            c.number_format = '0.00'
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        for c in row:
            c.number_format = '0'
    widths = [26, 15, 12, 15, 14, 8, 20, 14, 8]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[col_letter(i)].width = w
    ws.freeze_panes = 'A2'
    if ws.max_row >= 2:
        ws.auto_filter.ref = 'A1:' + col_letter(len(headers)) + str(ws.max_row)
    wb.save(out_path)
    return len(rows)


def build_full_xlsx(data, out_path):
    from openpyxl import Workbook
    headers = data.get('fullHeaders') or []
    rows = data.get('fullRows') or []
    wb = Workbook()
    ws = wb.active
    ws.title = '会员全部字段'
    ws.append(headers)
    style_header(ws, len(headers))
    for r in rows:
        ws.append(r)
    text_cols = []
    money_cols = []
    for i, h in enumerate(headers, 1):
        if '手机号' in h or '电话' in h:
            text_cols.append(i)
        if '·分' in h:
            money_cols.append(i)
    for col in text_cols:
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            for c in row:
                if c.value is not None:
                    c.value = str(c.value)
                c.number_format = '@'
    for col in money_cols:
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            for c in row:
                c.number_format = '#,##0'
    for i in range(1, len(headers) + 1):
        h = headers[i - 1]
        w = 16
        if i <= 3:
            w = 22
        elif '时间' in h or '日期' in h:
            w = 20
        elif '地址' in h or '备注' in h or '门店' in h:
            w = 26
        ws.column_dimensions[col_letter(i)].width = w
    ws.freeze_panes = 'A2'
    if ws.max_row >= 2:
        ws.auto_filter.ref = 'A1:' + col_letter(max(len(headers), 1)) + str(ws.max_row)
    wb.save(out_path)
    return len(rows)


def g(d, *ks):
    cur = d
    for k in ks:
        if not isinstance(cur, dict):
            return None
        cur = cur.get(k)
    return cur


def build_detail_xlsx(deep, out_path):
    from openpyxl import Workbook
    details = deep.get('details') or []
    wb = Workbook()

    # Sheet1 详情汇总
    ws = wb.active
    ws.title = '详情汇总'
    headers = ['会员ID', '姓名', '手机号', '级别', '标签', '充值余额·分', '赠送余额·分', '积分', '积分金额·分',
               '券数量', '券总额·分', 'VIP卡数量', '订单笔数', '订单总额·分', '订单件数', '最近订单时间',
               '最近充值时间', '最近充值金额·分', '最近充值门店', '到期日期']
    ws.append(headers)
    style_header(ws, len(headers))
    order_rows = []
    card_rows = []
    for d in details:
        uid = d.get('uid')
        name = d.get('name') or ''
        phone = d.get('phone') or ''
        order = d.get('order') or {}
        olist = order.get('list') if isinstance(order, dict) else None
        latest_order_time = ''
        if isinstance(olist, list) and olist:
            ts = [o.get('ctime') or '' for o in olist if isinstance(o, dict)]
            ts = [t for t in ts if t]
            if ts:
                latest_order_time = max(ts)
            for o in olist:
                if isinstance(o, dict):
                    order_rows.append([uid, name, phone, o.get('orderid'), o.get('sncode'),
                                       o.get('trmb'), o.get('wstatus'), o.get('ctime')])
        vipcard = d.get('vipcard') or {}
        vlist = vipcard.get('list') if isinstance(vipcard, dict) else None
        vipcard_n = len(vlist) if isinstance(vlist, list) else 0
        if isinstance(vlist, list):
            for it in vlist:
                card_rows.append(['VIP卡', uid, name, phone, json.dumps(it, ensure_ascii=False) if isinstance(it, (dict, list)) else it])
        ticket = d.get('ticket') or {}
        tlist = ticket.get('list') if isinstance(ticket, dict) else None
        if isinstance(tlist, list):
            for it in tlist:
                card_rows.append(['优惠券', uid, name, phone, json.dumps(it, ensure_ascii=False) if isinstance(it, (dict, list)) else it])
        ws.append([
            uid, name, phone, d.get('vip_name') or d.get('vipname') or '', d.get('tagname') or '',
            d.get('cbalance'), d.get('zbalance'), d.get('score'), d.get('scorermb'),
            (ticket.get('num') if isinstance(ticket, dict) else None), (ticket.get('rmb') if isinstance(ticket, dict) else None),
            vipcard_n, order.get('cnt') if isinstance(order, dict) else None,
            order.get('trmb') if isinstance(order, dict) else None, order.get('wcnt') if isinstance(order, dict) else None,
            latest_order_time, d.get('chargetime'), d.get('chargermb'), d.get('chargeshop'), d.get('edate'),
        ])
    for col in (3,):
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            for c in row:
                if c.value is not None:
                    c.value = str(c.value)
                c.number_format = '@'
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[col_letter(i)].width = 22 if i <= 3 else (14 if '·分' not in h and '时间' not in h else 20)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = 'A1:' + col_letter(len(headers)) + str(max(ws.max_row, 1))

    if order_rows:
        ws2 = wb.create_sheet('订单明细')
        oh = ['会员ID', '姓名', '手机号', '订单号', '业务码', '订单金额·分', '状态码', '下单时间']
        ws2.append(oh)
        style_header(ws2, len(oh))
        for r in order_rows:
            ws2.append(r)
        for col in (3,):
            for row in ws2.iter_rows(min_row=2, min_col=col, max_col=col):
                for c in row:
                    if c.value is not None:
                        c.value = str(c.value)
                    c.number_format = '@'
        widths2 = [14, 18, 15, 14, 12, 14, 10, 20]
        for i, w in enumerate(widths2, 1):
            ws2.column_dimensions[col_letter(i)].width = w
        ws2.freeze_panes = 'A2'

    if card_rows:
        ws3 = wb.create_sheet('卡券明细')
        ch = ['类别', '会员ID', '姓名', '手机号', '内容']
        ws3.append(ch)
        style_header(ws3, len(ch))
        for r in card_rows:
            ws3.append(r)
        for col in (4,):
            for row in ws3.iter_rows(min_row=2, min_col=col, max_col=col):
                for c in row:
                    if c.value is not None:
                        c.value = str(c.value)
                    c.number_format = '@'
        widths3 = [10, 14, 18, 15, 80]
        for i, w in enumerate(widths3, 1):
            ws3.column_dimensions[col_letter(i)].width = w
        ws3.freeze_panes = 'A2'

    wb.save(out_path)
    return len(details)


def main():
    base = os.path.dirname(os.path.abspath(__file__))
    root = os.path.dirname(base)
    out_dir = os.path.join(root, '导出结果')
    ptr = os.path.join(out_dir, '.last_export.json')

    if not os.path.exists(ptr):
        print('[错误] 找不到 .last_export.json，请先运行导出脚本。')
        return 1

    with open(ptr, encoding='utf-8') as f:
        meta = json.load(f)

    json_path = meta.get('json') or ''
    if not json_path or not os.path.exists(json_path):
        print('[错误] 找不到数据文件：', json_path)
        return 1

    with open(json_path, encoding='utf-8') as f:
        data = json.load(f)

    rc = 0
    try:
        n = build_summary_xlsx(data, meta.get('xlsx') or (json_path[:-5] + '.xlsx'))
        print('Excel 已生成（常规表）：', meta.get('xlsx'))
        print('（共 %d 行数据）' % n)
    except Exception as e:
        print('[提示] 常规表生成失败：', e)
        rc = 1

    try:
        full_xlsx = meta.get('fullXlsx') or (meta.get('fullCsv') or '').replace('.csv', '.xlsx')
        n = build_full_xlsx(data, full_xlsx)
        print('Excel 已生成（全量字段）：', full_xlsx)
        print('（共 %d 行 × %d 列）' % (n, len(data.get('fullHeaders') or [])))
    except Exception as e:
        print('[提示] 全量字段表生成失败：', e)
        rc = 1

    deep_json = meta.get('deepJson')
    if deep_json and os.path.exists(deep_json):
        try:
            with open(deep_json, encoding='utf-8') as f:
                deep = json.load(f)
            deep_xlsx = meta.get('deepXlsx') or deep_json.replace('.json', '.xlsx')
            n = build_detail_xlsx(deep, deep_xlsx)
            print('Excel 已生成（会员详情）：', deep_xlsx)
            print('（共 %d 个会员）' % n)
        except Exception as e:
            print('[提示] 会员详情表生成失败：', e)
            rc = 1

    return rc


if __name__ == '__main__':
    sys.exit(main())
